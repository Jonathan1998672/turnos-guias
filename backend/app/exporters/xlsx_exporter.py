import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.exporters.base import RolExportable

AZUL_OSCURO = "FF1E293B"
GRIS_CLARO = "FFF2F2F2"
BORDE = Border(*[Side(style="thin", color="FF999999")] * 4)


class XlsxExportador:
    extension = "xlsx"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def exportar(self, rol: RolExportable) -> bytes:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Rol"

        total_columnas = len(rol.encabezados)

        hoja.append([rol.titulo])
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columnas)
        hoja["A1"].font = Font(bold=True, size=14)
        hoja["A1"].alignment = Alignment(horizontal="center")

        hoja.append([f"Turno: {rol.turno}"])
        hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columnas)
        hoja["A2"].alignment = Alignment(horizontal="center")

        hoja.append([])
        fila_encabezado = hoja.max_row + 1

        hoja.append(rol.encabezados)
        for celda in hoja[fila_encabezado]:
            celda.font = Font(bold=True, color="FFFFFFFF")
            celda.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
            celda.border = BORDE
            celda.alignment = Alignment(horizontal="center", vertical="center")

        for fila in rol.como_tabla()[1:]:
            hoja.append(fila)
            for indice, celda in enumerate(hoja[hoja.max_row]):
                celda.border = BORDE
                if indice == total_columnas - 1:
                    celda.font = Font(bold=True)
                    celda.fill = PatternFill("solid", fgColor=GRIS_CLARO)
                    celda.alignment = Alignment(horizontal="center")

        if rol.notas:
            hoja.append([])
            hoja.append([f"Notas: {rol.notas}"])

        self._ajustar_anchos(hoja, rol)
        hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=2)

        buffer = io.BytesIO()
        libro.save(buffer)
        return buffer.getvalue()

    def _ajustar_anchos(self, hoja, rol: RolExportable) -> None:
        tabla = rol.como_tabla()

        for columna in range(1, len(rol.encabezados) + 1):
            mas_largo = max(len(str(fila[columna - 1])) for fila in tabla)
            hoja.column_dimensions[get_column_letter(columna)].width = min(mas_largo + 4, 34)
